# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.8); s.right_margin = Cm(2.8)

def setf(run, name='宋体', size=11, bold=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.font.bold = bold

def title(text, size=18):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setf(p.add_run(text), '黑体', size, True)

def h(text, size=14):
    p = doc.add_paragraph()
    setf(p.add_run(text), '黑体', size, True)
    p.paragraph_format.space_before = Pt(10)

def body(text, indent=True):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.5
    setf(p.add_run(text), '宋体', 11)

title('小组分工与汇报叙述稿')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
setf(p.add_run('第一阶段实验 · 软件缺陷的发现与分类'), '楷体', 12)
doc.add_paragraph()

h('一、小组分工')
t = doc.add_table(rows=4, cols=3); t.style = 'Light Grid Accent 1'
data = [
    ('成员', '角色', '主要负责内容'),
    ('同学 A', '组长', '1) system.c / admin.c 静态代码审查；2) 登录、用户管理、删除学生路径的动态测试；3) 撰写缺陷报告中 BUG-001 ~ BUG-009；4) PPT 第 1-6 页制作与汇报'),
    ('同学 B', '组员', '1) teacher.c / student.c 静态代码审查；2) 成绩录入、修改、查询、排名路径的动态测试；3) 撰写缺陷报告中 BUG-010 ~ BUG-018；4) PPT 第 7-10 页制作与汇报'),
    ('共同', '—', '缺陷分类定级讨论、报告统稿、PPT 彩排与答辩准备'),
]
for i, row in enumerate(data):
    for j, v in enumerate(row):
        c = t.cell(i, j).paragraphs[0]
        setf(c.add_run(v), '宋体', 10.5, bold=(i == 0))

h('二、汇报叙述稿（约 4 分钟）')

body('【同学 A 开场，约 2 分钟】')
body('各位老师、同学好，我们小组汇报的是学生成绩管理系统第一阶段——软件缺陷的发现与分类。')
body('我们的被测对象是一个 C 语言控制台版的学生成绩管理系统，分管理员、教师、学生三类角色，'
     '支持登录、学生增删、成绩录入修改、查询排名等功能，数据存在 users.dat 和 students.dat 两个二进制文件里，共约 1100 行代码。')
body('测试方法上，我们一方面通读了全部源码做静态审查，重点看输入校验、数组越界、数据一致性和权限；'
     '另一方面实际运行程序，模拟正常和异常输入做动态测试。我们把缺陷按严重程度分成四级：致命、严重、一般、轻微。')
body('这一轮我们一共发现了 18 个缺陷，其中致命 4 个、严重 8 个、一般 3 个、轻微 3 个。'
     '问题主要集中在三块：数据一致性、安全防护和异常处理。')

body('【同学 B 接典型案例与总结，约 2 分钟】')
body('我举两个最典型的例子。第一个是 BUG-001，删除学生时会误删下一个学生的登录账号。'
     '原因是代码先把学生数组前移覆盖了被删位置，再用这个已经被覆盖的位置去匹配用户表，结果匹配到了下一个学生。')
body('第二个是安全问题：所有密码明文存在 dat 文件里，管理员在"查看所有用户"时界面直接打印明文密码，'
     '而且默认账号就是 admin/admin。此外还有重名学生登录串号、成绩重复录入导致平均分算错、'
     '超长输入后缓冲区残留导致后续输入串台等问题。')
body('针对这些问题，我们建议优先修复 4 个致命缺陷，然后补齐输入校验和成绩联动逻辑。'
     '下一阶段，我们会针对 deleteStudent、addScores 这些核心函数做白盒测试和代码覆盖率分析。')
body('我们的汇报到此结束，请老师和同学批评指正。')

h('三、可能被问到的问题（备用）')
body('Q：你们为什么不用专业测试工具？ A：本阶段以缺陷发现和分类为主，采用人工静态审查加黑盒执行，白盒工具会在第二阶段引入。', indent=False)
body('Q：致命缺陷是怎么定的？ A：按"是否导致数据错误/崩溃/安全漏洞"判定，例如删错账号属于数据级错误，定为致命。', indent=False)
body('Q：你们自己写的系统为什么还测？ A：课程要求就是先开发再测试，故意保留的实现缺陷正好用于练习发现和分类。', indent=False)

doc.save(r'F:\study\软件质量保证与测试\小组分工与汇报稿.docx')
print('OK')
