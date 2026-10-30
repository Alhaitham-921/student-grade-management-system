from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
OLD = ROOT / '修复前备份'
NEW = ROOT / 'zuoyezuihou'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.line_spacing = Pt(15)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):
    sec._sectPr.remove(grid)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)

def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text, 2)
def page(title):
    heading=doc.add_heading(title, 1)
    heading.paragraph_format.page_break_before=True
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row): c.text=str(v)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3' if i==0 else 'FFFFFF');tcPr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
            tcPr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.line_spacing=1.08
                for r in para.runs: r.font.size=Pt(9.5);r.bold=i==0
    p('')

def loc(file, needle, old=False):
    path=(OLD if old else NEW)/file
    text=path.read_bytes().decode('gbk' if old and path.suffix in ['.c','.h'] else 'utf-8-sig')
    lines=text.splitlines()
    for i,line in enumerate(lines,1):
        if needle in line: return f'{file} 第{i}行'
    raise ValueError((file,needle))
def code(label, file, needle, old=False, count=1):
    path=(OLD if old else NEW)/file
    lines=path.read_bytes().decode('gbk' if old and path.suffix in ['.c','.h'] else 'utf-8-sig').splitlines()
    n=next(i for i,line in enumerate(lines) if needle in line)
    para=p(f'{label}  {loc(file,needle,old)}')
    para.paragraph_format.keep_with_next=True
    for line in lines[n:n+count]:
        para=doc.add_paragraph()
        para.paragraph_format.space_after=Pt(2)
        para.paragraph_format.line_spacing=Pt(12)
        r=para.add_run(line.strip());r.font.name='Consolas';r.font.size=Pt(9)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')

doc.add_heading('学生成绩管理系统修复前后对比',0)
p('代码修改说明与手动验证参考    2026年10月29日')
p('本说明对比修复前备份与当前工程，解释五个 C 源文件及相关头文件、VS2022 配置的实际改动。重点是哪些输入或操作原来会出错、现在如何处理，以及怎样通过程序界面观察差异。')
table(['文件','主要修改'],[
('main.c','设置 UTF-8 控制台；加载和初始化后保存迁移结果；统一清屏入口'),
('system.c','完整行输入、严格数字解析、密码哈希、登录冷却、数据校验与迁移、平均分计算、文件保存'),
('admin.c','学生按学号建账号；容量检查；删除账号关联修复；密码修改和隐藏'),
('teacher.c','阻止重复课程；满课程不再清空；空成绩提前提示；排名分支和零分显示修复'),
('student.c','无授课教师提示；统一界面暂停和清屏'),
('system.h 及工程配置','统一声明和数组上限；扩展密码字段；设置 /utf-8 和调试工作目录')],[1.65,5.35])
h('对比依据')
p('修改前：工作区“修复前备份”目录内的源码。修改后：“zuoyezuihou”目录内的当前源码。文中代码均摘自这两套文件，行号对应本次整理时的文件版本。')
p('“修复版手工测试”是独立工程副本。该副本建立时没有业务数据，账号需要按测试流程自行创建；后续测试输入会保存在副本目录。')
h('验证口径')
p('前期已完成 VS2022 下 Debug、Release 与 x64、x86 四种配置编译，并执行核心自动回归。本说明中的手动观察项用于你自行复核，不代表已替你执行全部界面测试，也不代表原缺陷报告所有问题都已彻底解决。')

page('一 输入检查与字母输入的处理')
h('修改位置')
p(loc('system.c','static void readLine')+'；'+loc('system.c','int safeReadInt')+'；'+loc('system.c','float safeReadFloat')+'。')
code('修改前','system.c','if (sscanf(buffer, "%d", &value)',True)
code('现在','system.c','long value = strtol',count=4)
p('原来 sscanf 只要解析到开头的数字就算成功，未核对整行是否读完。现在使用 strtol、strtof，检查结束位置、转换错误和取值范围；分数还要求是有限数。末尾仅有空白时允许通过。')
table(['输入场景','修改前','现在'],[
('菜单输入 abc','转换失败，要求重输','仍然拒绝；这不是本次新增能力'),
('菜单输入 1abc','可按选项1接受','拒绝，要求完整整数'),
('成绩输入 80abc','可按80分接受','拒绝，不保存'),
('成绩输入 -1 或101','已有0至100范围检查','保留范围检查'),
('字符串直接回车或仅空格','可能作为有效字段保存','重新输入非空内容'),
('用户名超过19字节','余下字符可能被下一项读取','丢弃超长行，重新输入当前字段')],[1.8,2.5,2.7])
h('字符串现在怎样读取')
p('readLine 读取完整一行；超长时清理该行剩余内容。safeReadString 在其上增加非空检查。提示输出由 printf(prompt) 改为 printf("%s", prompt)，避免提示中的百分号被当成格式指令。输入流结束时退出，避免反复读 EOF。')
h('没有新增的限制')
p('用户名、姓名、密码仍可包含字母。学号目前也没有“只能是数字”的规则，abc 形式的学号仍可创建。长度按字节计算；UTF-8 中文通常占多个字节，不能把19字节理解为19个汉字。')

page('二 学生账号与删除关联')
h('学生账号改用学号')
code('修改前','admin.c','strcpy(studentUser.username, newStudent.name)',True)
code('现在','admin.c','strcpy(studentUser.username, newStudent.stuNum)')
code('修改前的查档条件','system.c','strcmp(students[i].name, username)',True)
code('现在的查档条件','system.c','strcmp(students[i].stuNum, username)')
p('原来两位同名学生可能共用姓名标识，登录后查找总落到第一份档案。现在账号创建、登录后的档案查找、删除账号都以学号关联；新建学号同时检查学生档案和用户账号是否已占用。初始密码仍是学号，但保存为哈希。')
h('删除前先保存目标学号')
code('现在','admin.c','char deletedId[STU_NUM_LEN]',count=2)
code('现在的账号匹配条件','admin.c','users[i].type == STUDENT && strcmp(users[i].username, deletedId)')
p('原来先把 students 数组前移，再读取 students[index].name 去找账号。此时该位置可能已属于下一名学生，导致删错账号。现在先保存被删学生学号，再移动数组，并只删除类型为 STUDENT、用户名等于目标学号的账号。')
h('容量不足时提前拒绝')
code('现在','admin.c','if (studentCount >= MAX_STUDENTS || userCount >= MAX_USERS)',count=3)
p('原来学生档案先入数组，用户满额后只提示未建账号，形成“有档案不能登录”。现在在读取新增信息前检查两种容量；在当前单进程交互流程中避免该中间状态。这不是两个数据文件之间的事务保证。')
h('手动观察')
p('建立学号2026001、2026003的两位张三，分别录入不同成绩；各自登录应只见本人档案。另建李四2026002，删除2026001后，李四账号仍应能登录。')

page('三 成绩录入与统计结果')
h('同一课程不能重复添加')
code('现在','teacher.c','if (findCourseIndex(student, currentUser.teachCourse) >= 0)',count=3)
p('原来每次都追加当前教师所授课程，且允许一次添加多条同名课程。现在每次只添加当前课程的一条成绩；已有该课时提示使用“修改成绩”，不会重复累加。重复检查在满额检查之前执行。')
h('课程满额不再清空已有记录')
code('修改前','teacher.c','students[studentIndex].courseCount >= MAX_COURSES)',True,count=3)
code('现在','teacher.c','if (student->courseCount >= MAX_COURSES)',count=3)
p('旧版把课程数达到10也当成异常，重置 courseCount 为0，后续录入覆盖原成绩。现在达到上限直接拒绝添加。加载阶段另外检查课程数是否落在0至10之间。')
h('平均分在保存时统一更新')
code('现在','system.c','void saveAll(void)',count=4)
p('原来 addScores、modifyScores 保存成绩时没有更新 average，学生端和排名可能显示旧值。现在 saveAll 先调用 refreshAverages，逐个学生按现有课程计算均值；无课程时记为0。加载数据后也会重新计算。')
p('观察例子：数学80、英语60时，直接查看应为70；将数学改为100后，不手动执行“计算平均分”，再次查看应为80。排名仍按学生全部已录课程的平均分进行。')
h('零分科目与无成绩提示')
p(loc('teacher.c','if (k == 0 || tempStudents[i].courses[k].score > maxScore)')+'：最高分判断包含第一门课程，即使全为0分，也能显示科目名称。')
p(loc('teacher.c','if (students[studentIndex].courseCount == 0)')+'：修改成绩前先检查是否有成绩，避免先打印空表。后面的旧检查仍保留，属于冗余判断，不影响这一行为。')

page('四 密码存储与登录控制')
h('明文密码改为带盐哈希')
code('修改前','system.c','strcmp(users[i].password, password) == 0',True)
code('现在','system.c','verifyPassword(password, users[i].password)')
p('新增 derivePassword、hashPassword、verifyPassword。使用 Windows BCrypt，生成16字节随机盐，执行100000次 PBKDF2-HMAC-SHA256，得到32字节派生值。保存格式为 p1$ 加盐和派生值的十六进制文本，共99个字符；User.password 从20扩为128字节。验证时重新派生并比较结果。')
table(['位置','原来','现在'],[
('管理员查看用户','直接显示每个人的密码','固定显示 ******'),
('创建教师和学生','保存原始密码字符串','保存带盐哈希'),
('修改密码','比较并覆盖明文','验证旧密码哈希，再保存新哈希'),
('清空所有用户','管理员密码重置为admin','保留当前管理员密码的哈希')],[1.5,2.45,3.05])
h('默认管理员必须改密码')
p('初始化仍建立 admin / admin，便于首次运行。登录发现管理员密码仍为 admin 时，必须完成修改，才能进入管理员菜单。修改密码要求非空、至少6字节且两次一致；最小长度只用于修改密码，创建教师时尚未统一执行6字节限制。')
h('三次失败后暂时冷却')
code('现在','system.c','static time_t lockedUntil = 0;',count=5)
p('原来 int attempts = 3 每次进入登录函数都会重新计数。现在使用 LOGIN_ATTEMPTS 宏，耗尽一次登录流程的尝试次数后，设置 lockedUntil = time(NULL) + 60；期间再次进入登录会被拒绝。')
p('当前冷却是进程内全局状态，不区分账号或角色，重启会清除；并未实现按账号持久锁定。用户输入密码时仍会回显，“用户列表隐藏密码”不等于“输入时星号遮罩”。')

page('五 数据文件校验与旧数据迁移')
h('读取前验证记录数量和文件长度')
code('修改前','system.c','fread(&userCount, sizeof(int), 1, file)',True,count=2)
code('现在','system.c','if (fread(&n, sizeof(n), 1, file)',count=3)
p('旧版直接使用文件中的数量读取数组，数量越界或文件截断可能导致异常。现在 loadRecords 检查数量上下限、实际读到的记录数、文件尾及读取错误。loadAll 继续检查字符串结束符、角色范围、课程数、教师数及分数范围。异常时报告并退出，避免按错误数量访问数组。')
h('兼容原始用户记录')
p('User 密码字段扩展导致记录大小变化。加载时按“4字节计数＋76字节×记录数”识别原始用户文件，再读取旧结构并将明文密码转为哈希。旧格式相关文字从GBK转换为UTF-8；转换失败或目标字段放不下时拒绝加载，不静默截断。')
h('修复旧账号与学生档案之间的关联')
p('能唯一匹配学生姓名的旧学生账号改为学号；同名无法唯一匹配或用户名冲突时停止加载，要求人工恢复。无对应档案的账号仍保留在用户数组中，但登录时因找不到档案而被拒绝。学生档案缺失账号时，若学号未被占用且用户容量足够，恢复学号账号，初始密码为学号。')
p('原始样本的账号1与档案不匹配，学号1234对应账号缺失。该恢复规则用于处理这些旧数据；空数据手工测试副本不需要执行这类迁移。')
h('先写临时文件再替换单个文件')
code('现在','system.c','if (!ok || !MoveFileExA',count=3)
p('旧版直接以 wb 打开原文件，保存失败可能留下截断数据。现在先写 .tmp 文件，检查写入和关闭结果，再通过 MoveFileExA 替换目标文件。main 在加载、初始化成功后调用 saveAll，将迁移结果保存为新格式。')
p('限制：users.dat 和 students.dat 分别替换，缺少跨文件事务；两者之间失败仍可能造成不一致。新数据文件不能交给旧程序读取，恢复旧版需要同时恢复程序和配套数据。')

page('六 VS2022 配置与界面调整')
table(['修改处','修改前','现在的实现'],[
('safeReadString 声明','system.h 使用int，admin.c另写size_t声明','统一为size_t，移除admin.c重复extern声明'),
('学生数组声明','头文件写MAX_USERS，定义用MAX_STUDENTS','头文件统一为MAX_STUDENTS'),
('排名switch分支','case后直接声明局部变量','相关case增加花括号，明确局部作用域'),
('源码与控制台编码','源码GBK，未主动设置控制台编码','源码UTF-8；/utf-8；设置输入输出代码页'),
('编译配置','旧CRT接口可能产生C4996诊断','项目定义_CRT_SECURE_NO_WARNINGS'),
('数据目录','依赖启动时所在目录','VS调试工作目录指定为$(ProjectDir)'),
('清屏与暂停','各菜单直接调用cls和pause','统一clearScreen、pauseScreen；交互终端按回车继续'),
('学生无教师记录','教师标题后没有明确说明','teacherCount为0时显示暂无授课教师')],[1.45,2.5,3.05])
h('相关代码位置')
p(loc('main.c','SetConsoleCP(CP_UTF8)')+'：设置控制台输入与输出为UTF-8。'+loc('system.h','void safeReadString')+'：统一读取函数的参数类型。')
p(loc('system.c','void clearScreen(void)')+'：通过 _isatty 判断交互终端，重定向输入时跳过清屏。pauseScreen 同样只在交互输入时暂停。')
h('实现边界')
p('当前目标仍是 Windows 与 VS2022。程序仍依赖 windows.h、io.h、BCrypt 和 cls，没有完成 Linux/macOS 兼容。抑制旧CRT接口诊断也不等于所有接口都已替换为安全版本。')
p('部分旧注释尚未同步，例如 addStudent 注释仍写“用户名=姓名”，实际赋值已经使用学号。判断实现应以执行代码为准，后续可以单独清理注释。')

page('七 手动验证与当前限制')
h('独立测试账号')
p('打开“修复版手工测试”里的工程或“开始手动测试.bat”。首次管理员为admin / admin；完成首次改密后使用admin / admin123。下列教师与学生需要自行创建，并非副本预置账号。')
table(['角色','用户名','测试密码','设置'],[
('数学教师','teacher1','teacher123','授课科目数学'),
('英语教师','teacher2','teacher123','授课科目英语'),
('学生张三','2026001','2026001','分配数学和英语教师'),
('学生李四','2026002','2026002','分配数学教师'),
('同名张三','2026003','2026003','分配数学教师')],[1.2,1.4,1.5,2.9])
h('优先观察的六组结果')
for text in [
'1  输入菜单1abc、成绩80abc：新版均拒绝；输入纯abc，旧版本来就应拒绝。',
'2  给2026001添加数学80、英语60：无需手工计算，平均分应为70；数学改100后应为80。',
'3  再给同一学生添加数学：提示课程已存在；另一位张三2026003应只能看到自己的成绩。',
'4  删除2026001：2026002仍能登录，且原有成绩保留。',
'5  查看用户列表：密码显示******；修改密码后退出重开，新密码可登录、旧密码失败。',
'6  连续三次错误登录：立即重试会被暂时拒绝，等待至少60秒后恢复。']:
    p(text)
h('不能据此宣称全部解决的部分')
p('学号未限定纯数字；密码规则未统一覆盖所有创建入口；密码输入仍回显；登录冷却不持久化；两个数据文件无跨文件事务；跨平台兼容未完成；历史重复课程不会自动合并。')
p('课程满额、用户满额、无授课教师等边界需要专门准备数据，不能只凭普通菜单操作就视为已人工验证。此次文档整理没有代替你执行手动测试，也没有启动测试程序或更改业务数据。')

out=ROOT/'学生成绩管理系统_修改前后实现对比.docx'
doc.save(out)
print(out)
